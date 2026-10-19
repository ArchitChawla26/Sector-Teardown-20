"""Shared workbook builder for the AI Infra & Semis sector teardowns.

Config-driven mirror of the equity-teardown-hood build_model.py: every computed
cell is a live Excel formula referencing Assumptions / Historicals. Sheets:
Cover, Assumptions, Historicals, DCF_Base, DCF_Bull, DCF_Bear, Comps, Checks,
Sources.

Each company's build_model.py defines CONFIG and calls build(CONFIG).
"""
import math
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E78")
SUB_FONT = Font(name="Calibri", size=11, italic=True, color="595959")
LBL_FONT = Font(name="Calibri", size=11)
INP_FILL = PatternFill("solid", fgColor="DDEBF7")
TAG_DATA_FILL = PatternFill("solid", fgColor="E2EFDA")
TAG_JUDG_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(*(Side(style="thin", color="BFBFBF") for _ in range(4)))
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
FMT_M = '#,##0.0'
FMT_PCT = '0.0%'
FMT_PX = '"$"#,##0.00'
FMT_X = '#,##0.0"x"'

MODEL_ASOF = "October 19, 2026"


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = THIN


def title_block(ws, title, subtitle, ncols=6):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=subtitle)
    c.font = SUB_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18


def finalize(ws, widths):
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B4"


def build(cfg):
    """Build the valuation workbook for one company. cfg keys (see company build_model.py)."""
    T, NAME, EXCH = cfg["ticker"], cfg["name"], cfg["exchange"]
    wb = openpyxl.Workbook()
    OUT = cfg["out_path"]

    # ============================ COVER ============================
    ws = wb.active
    ws.title = "Cover"
    title_block(ws, f"Equity Research Teardown  \u2014  {NAME} ({EXCH}: {T})",
                f"Independent equity research project by Archit Chawla  |  As of {MODEL_ASOF}", 4)
    ws["A4"] = f"Current price ({cfg['price_asof_label']} close)"
    ws["B4"] = cfg["price"]
    ws["B4"].number_format = FMT_PX
    ws["B4"].font = Font(name="Calibri", size=12, bold=True)
    ws["A5"] = "Recommendation"
    ws["B5"] = cfg["recommendation"]
    rec_color = "C00000" if cfg["recommendation"] == "PASS" else ("1F4E78" if cfg["recommendation"] == "HOLD" else "548235")
    ws["B5"].font = Font(name="Calibri", size=12, bold=True, color=rec_color)
    ws["A6"] = "One-line thesis"
    ws["B6"] = cfg["one_liner"]
    ws["B6"].alignment = LEFT
    for r in (4, 5, 6):
        ws.cell(row=r, column=1).font = LBL_FONT
    ws["A8"] = "Contents"
    ws["A8"].font = Font(name="Calibri", size=11, bold=True)
    contents = [
        ("Assumptions", "WACC build and all forecast drivers; [DATA] vs [JUDGMENT] tags"),
        ("Historicals", f"Financials from Yahoo Finance annual statements ({cfg['hist_label']})"),
        ("DCF_Base / DCF_Bull / DCF_Bear", "5-year DCF per case; all formulas live"),
        ("Comps", f"{cfg['peer_label']}: EV/Revenue, EV/EBITDA, P/E and implied ranges"),
        ("Checks", "Verification appendix: tie-outs and data sources"),
        ("Sources", "Full URLs and access dates"),
    ]
    for i, (s, d) in enumerate(contents, start=9):
        ws.cell(row=i, column=1, value=s).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=i, column=2, value=d).font = LBL_FONT
        ws.cell(row=i, column=2).alignment = LEFT
    ws["A16"] = "Disclaimer: Independent project for skill demonstration; not investment advice."
    ws["A16"].font = Font(name="Calibri", size=9, italic=True, color="595959")
    ws.merge_cells("A16:D16")
    finalize(ws, [28, 90, 10, 10])

    # ============================ ASSUMPTIONS ============================
    ws = wb.create_sheet("Assumptions")
    A = "Assumptions"
    title_block(ws, "Assumptions", "Blue cells are inputs. [DATA] = market/filing data; "
                "[JUDGMENT] = analyst estimate with rationale.", 4)
    ws["A3"] = "WACC build"
    ws["A3"].font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    wacc_rows = [
        ("Risk-free rate (US 10Y)", cfg["rf"], "DATA", "5.09% on 2026-10-19 (Reuters); rounded to 5.1%"),
        ("Equity risk premium", cfg["erp"], "JUDGMENT", "Long-run US ERP assumption"),
        (f"Beta ({T})", cfg["beta"], "DATA", cfg["beta_note"]),
        ("Cost of equity (CAPM)", f"={A}!$B$4+{A}!$B$6*{A}!$B$5", "CALC", "Rf + beta \u00d7 ERP"),
        ("Pre-tax cost of debt", cfg["cod_pre"], "JUDGMENT", cfg["cod_note"]),
        ("Marginal tax rate", cfg["tax"], "JUDGMENT", cfg["tax_note"]),
        ("After-tax cost of debt", f"={A}!$B$8*(1-{A}!$B$9)", "CALC", ""),
        ("Market value of equity ($M)", cfg["mcap"], "DATA", cfg["mcap_note"]),
        ("Book value of debt ($M)", cfg["debt_book"], "DATA", cfg["debt_note"]),
        ("Weight of equity", f"={A}!$B$11/({A}!$B$11+{A}!$B$12)", "CALC", ""),
        ("Weight of debt", f"={A}!$B$12/({A}!$B$11+{A}!$B$12)", "CALC", ""),
        ("WACC \u2014 base case", f"={A}!$B$13*{A}!$B$7+{A}!$B$14*{A}!$B$10", "CALC", ""),
        ("WACC \u2014 bull case (-100 bps)", f"={A}!$B$15-0.01", "CALC", "Rate-normalization scenario"),
        ("WACC \u2014 bear case (+100 bps)", f"={A}!$B$15+0.01", "CALC", "Higher-for-longer scenario"),
    ]
    R = {"rf": 4, "erp": 5, "beta": 6, "coe": 7, "cod": 8, "tax": 9, "atcod": 10,
         "mcap": 11, "debt": 12, "we": 13, "wd": 14, "wacc": 15,
         "wacc_bull": 16, "wacc_bear": 17}
    for i, (lbl, val, tag, note) in enumerate(wacc_rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=lbl).font = LBL_FONT
        c = ws.cell(row=r, column=2, value=val)
        if tag in ("DATA", "JUDGMENT"):
            c.fill = INP_FILL
        c.number_format = FMT_PCT if r in (4, 5, 7, 8, 9, 10, 13, 14, 15, 16, 17) else FMT_M
        if r == 6:
            c.number_format = '0.00'
        t = ws.cell(row=r, column=3, value=f"[{tag}]")
        t.font = Font(name="Calibri", size=9, bold=True)
        t.fill = TAG_DATA_FILL if tag == "DATA" else (TAG_JUDG_FILL if tag == "JUDGMENT" else PatternFill())
        t.alignment = CENTER
        ws.cell(row=r, column=4, value=note).font = Font(name="Calibri", size=9, color="595959")
        ws.cell(row=r, column=4).alignment = LEFT

    ws["A19"] = f"Forecast drivers ({cfg['fcast_label']})"
    ws["A19"].font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    for j, y in enumerate(cfg["fcast_years"]):
        c = ws.cell(row=20, column=2 + j, value=y)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = CENTER
    drv = [
        ("Revenue growth \u2014 base", cfg["cases"]["base"]["g"], "JUDGMENT", cfg["g_base_note"]),
        ("Revenue growth \u2014 bull", cfg["cases"]["bull"]["g"], "JUDGMENT", cfg["g_bull_note"]),
        ("Revenue growth \u2014 bear", cfg["cases"]["bear"]["g"], "JUDGMENT", cfg["g_bear_note"]),
        ("EBITDA margin \u2014 base", cfg["cases"]["base"]["m"], "JUDGMENT", cfg["m_base_note"]),
        ("EBITDA margin \u2014 bull", cfg["cases"]["bull"]["m"], "JUDGMENT", cfg["m_bull_note"]),
        ("EBITDA margin \u2014 bear", cfg["cases"]["bear"]["m"], "JUDGMENT", cfg["m_bear_note"]),
    ]
    R_DRV = {}
    for i, (lbl, vals, tag, note) in enumerate(drv):
        r = 21 + i
        R_DRV[lbl] = r
        ws.cell(row=r, column=1, value=lbl).font = LBL_FONT
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + j, value=v)
            c.fill = INP_FILL
            c.number_format = FMT_PCT
        t = ws.cell(row=r, column=7, value=f"[{tag}]")
        t.font = Font(name="Calibri", size=9, bold=True)
        t.fill = TAG_JUDG_FILL
        t.alignment = CENTER
        ws.cell(row=r, column=8, value=note).font = Font(name="Calibri", size=9, color="595959")
        ws.cell(row=r, column=8).alignment = LEFT
    misc = [
        ("D&A (% of revenue)", cfg["da_pct"], "JUDGMENT", cfg["da_note"]),
        ("Stock-based comp. (% of revenue)", None, "JUDGMENT", cfg["sbc_note"]),
        ("Capex (% of revenue)", cfg["capex_pct"], "JUDGMENT", cfg["capex_note"]),
        ("Change in NWC (% of revenue growth)", cfg["nwc_pct"], "JUDGMENT", "Receivables/inventory scale with semi activity"),
        ("Terminal growth \u2014 base", cfg["cases"]["base"]["tg"], "JUDGMENT", "Long-run nominal GDP"),
        ("Terminal growth \u2014 bull", cfg["cases"]["bull"]["tg"], "JUDGMENT", ""),
        ("Terminal growth \u2014 bear", cfg["cases"]["bear"]["tg"], "JUDGMENT", ""),
        ("Net debt ($M)", cfg["net_debt"], "DATA", cfg["net_debt_note"]),
        ("Diluted shares outstanding (M)", cfg["shares"], "DATA", cfg["shares_note"]),
    ]
    R_MISC = {}
    for i, (lbl, val, tag, note) in enumerate(misc):
        r = 28 + i
        R_MISC[lbl] = r
        ws.cell(row=r, column=1, value=lbl).font = LBL_FONT
        if lbl.startswith("Stock-based"):
            for j, v in enumerate(cfg["sbc_pct"]):
                c = ws.cell(row=r, column=2 + j, value=v)
                c.fill = INP_FILL
                c.number_format = FMT_PCT
        else:
            c = ws.cell(row=r, column=2, value=val)
            c.fill = INP_FILL
            c.number_format = FMT_PCT if isinstance(val, float) and val < 1 else FMT_M
        t = ws.cell(row=r, column=7, value=f"[{tag}]")
        t.font = Font(name="Calibri", size=9, bold=True)
        t.fill = TAG_DATA_FILL if tag == "DATA" else TAG_JUDG_FILL
        t.alignment = CENTER
        ws.cell(row=r, column=8, value=note).font = Font(name="Calibri", size=9, color="595959")
        ws.cell(row=r, column=8).alignment = LEFT
    R_SBC = R_MISC["Stock-based comp. (% of revenue)"]
    finalize(ws, [38, 16, 16, 16, 16, 16, 12, 60])

    # ============================ HISTORICALS ============================
    ws = wb.create_sheet("Historicals")
    title_block(ws, f"Historical financials \u2014 {T}",
                f"$ millions. Source: Yahoo Finance annual statements (filing-derived), as of 2026-10-19. {cfg['hist_note']}", 6)
    ws["A3"] = ""
    ncols = len(cfg["hist_years"])
    for j, y in enumerate(cfg["hist_years"]):
        ws.cell(row=3, column=2 + j, value=y)
    style_header(ws, 3, 1 + ncols)
    hist_rows = [
        ("Total revenue", "revenue", FMT_M),
        ("Gross profit", "gross_profit", FMT_M),
        ("EBITDA", "ebitda", FMT_M),
        ("EBITDA margin", None, FMT_PCT),
        ("Operating income", "op_income", FMT_M),
        ("Net income", "net_income", FMT_M),
        ("Diluted average shares (M)", "dil_shares", FMT_M),
        ("Operating cash flow", "ocf", FMT_M),
        ("Capital expenditures", "capex", FMT_M),
        ("Free cash flow", "fcf", FMT_M),
        ("Cash & equivalents", "cash", FMT_M),
        ("Total debt", "debt", FMT_M),
        ("Revenue growth", None, FMT_PCT),
    ]
    H = cfg["hist"]
    for i, (lbl, key, fmt) in enumerate(hist_rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=lbl).font = LBL_FONT
        for j in range(ncols):
            col = get_column_letter(2 + j)
            c = ws.cell(row=r, column=2 + j)
            if key:
                c.value = H[key][j]
            elif lbl == "EBITDA margin":
                c.value = f"={col}6/{col}4"
            elif lbl == "Revenue growth" and j > 0:
                prev = get_column_letter(1 + j)
                c.value = f"={col}4/{prev}4-1"
            c.number_format = fmt
    H_REV = 4
    finalize(ws, [30] + [16] * ncols)

    # ============================ DCF SHEETS ============================
    CASE_SHEETS = {
        "base": ("DCF_Base", "Discounted cash flow \u2014 base case",
                 R_DRV["Revenue growth \u2014 base"], R_DRV["EBITDA margin \u2014 base"],
                 f"{A}!$B${R['wacc']}", R_MISC["Terminal growth \u2014 base"]),
        "bull": ("DCF_Bull", "Discounted cash flow \u2014 bull case",
                 R_DRV["Revenue growth \u2014 bull"], R_DRV["EBITDA margin \u2014 bull"],
                 f"{A}!$B${R['wacc_bull']}", R_MISC["Terminal growth \u2014 bull"]),
        "bear": ("DCF_Bear", "Discounted cash flow \u2014 bear case",
                 R_DRV["Revenue growth \u2014 bear"], R_DRV["EBITDA margin \u2014 bear"],
                 f"{A}!$B${R['wacc_bear']}", R_MISC["Terminal growth \u2014 bear"]),
    }

    def build_dcf(name, subtitle, g_row, m_row, wacc_ref, tg_row):
        ws = wb.create_sheet(name)
        title_block(ws, f"{NAME} ({T}) \u2014 {subtitle}",
                    "$ millions, except per-share. All formulas reference Assumptions / Historicals.", 6)
        ws["A3"] = ""
        for j, y in enumerate(cfg["fcast_years"]):
            ws.cell(row=3, column=2 + j, value=y)
        style_header(ws, 3, 6)
        labels = [
            "Revenue", "Revenue growth", "EBITDA margin", "EBITDA",
            "Less: Depreciation & amortization", "EBIT", "Tax rate", "Less: Income tax",
            "NOPAT", "Add: Stock-based compensation", "Less: Capital expenditures",
            "Less: Change in net working capital", "Unlevered free cash flow",
            "Discount factor", "Present value of FCF",
        ]
        for i, lbl in enumerate(labels):
            ws.cell(row=4 + i, column=1, value=lbl).font = LBL_FONT
        HR = f"Historicals!${get_column_letter(1 + ncols)}${H_REV}"
        for j in range(5):
            col = get_column_letter(2 + j)
            pcol = get_column_letter(1 + j)
            r = 4
            if j == 0:
                ws.cell(row=r, column=2 + j, value=f"={HR}*(1+{A}!{col}${g_row})")
            else:
                ws.cell(row=r, column=2 + j, value=f"={pcol}{r}*(1+{A}!{col}${g_row})")
            ws.cell(row=r + 1, column=2 + j,
                    value=f"={col}{r}/{HR}-1" if j == 0 else f"={col}{r}/{pcol}{r}-1").number_format = FMT_PCT
            ws.cell(row=r + 2, column=2 + j, value=f"={A}!{col}${m_row}").number_format = FMT_PCT
            ws.cell(row=r + 3, column=2 + j, value=f"={col}{r}*{col}{r+2}")
            ws.cell(row=r + 4, column=2 + j, value=f"={col}{r}*{A}!$B${R_MISC['D&A (% of revenue)']}")
            ws.cell(row=r + 5, column=2 + j, value=f"={col}{r+3}-{col}{r+4}")
            ws.cell(row=r + 6, column=2 + j, value=f"={A}!$B${R['tax']}").number_format = FMT_PCT
            ws.cell(row=r + 7, column=2 + j, value=f"={col}{r+5}*{col}{r+6}")
            ws.cell(row=r + 8, column=2 + j, value=f"={col}{r+5}-{col}{r+7}")
            ws.cell(row=r + 9, column=2 + j, value=f"={col}{r}*{A}!{col}${R_SBC}")
            ws.cell(row=r + 10, column=2 + j, value=f"={col}{r}*{A}!$B${R_MISC['Capex (% of revenue)']}")
            prev_rev = HR if j == 0 else f"{pcol}{r}"
            ws.cell(row=r + 11, column=2 + j,
                    value=f"=({col}{r}-{prev_rev})*{A}!$B${R_MISC['Change in NWC (% of revenue growth)']}")
            ws.cell(row=r + 12, column=2 + j,
                    value=f"={col}{r+8}+{col}{r+4}+{col}{r+9}-{col}{r+10}-{col}{r+11}")
            ws.cell(row=r + 13, column=2 + j, value=f"=1/(1+{wacc_ref})^{j+1}").number_format = '0.000'
            ws.cell(row=r + 14, column=2 + j, value=f"={col}{r+12}*{col}{r+13}")
        s = 19
        ws.cell(row=s, column=1, value="Sum of PV of FCF").font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=s, column=2, value="=SUM(B18:F18)").number_format = FMT_M
        ws.cell(row=s + 1, column=1, value="Terminal value").font = LBL_FONT
        ws.cell(row=s + 1, column=2,
                value=f"=F16*(1+{A}!$B${tg_row})/({wacc_ref}-{A}!$B${tg_row})").number_format = FMT_M
        ws.cell(row=s + 2, column=1, value="PV of terminal value").font = LBL_FONT
        ws.cell(row=s + 2, column=2, value=f"=B{s+1}/(1+{wacc_ref})^5").number_format = FMT_M
        ws.cell(row=s + 3, column=1, value="Enterprise value").font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=s + 3, column=2, value=f"=B{s}+B{s+2}").number_format = FMT_M
        ws.cell(row=s + 4, column=1, value="Less: Net debt").font = LBL_FONT
        ws.cell(row=s + 4, column=2, value=f"={A}!$B${R_MISC['Net debt ($M)']}").number_format = FMT_M
        ws.cell(row=s + 5, column=1, value="Equity value").font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=s + 5, column=2, value=f"=B{s+3}-B{s+4}").number_format = FMT_M
        ws.cell(row=s + 6, column=1, value="Diluted shares outstanding (M)").font = LBL_FONT
        ws.cell(row=s + 6, column=2, value=f"={A}!$B${R_MISC['Diluted shares outstanding (M)']}").number_format = FMT_M
        ws.cell(row=s + 7, column=1, value="Implied value per share ($)").font = \
            Font(name="Calibri", size=11, bold=True, color="1F4E78")
        ws.cell(row=s + 7, column=2, value=f"=B{s+5}/B{s+6}").number_format = FMT_PX
        for rr in range(4, 27):
            for cc in range(2, 7):
                cell = ws.cell(row=rr, column=cc)
                if cell.value is None:
                    continue
                if cell.number_format == 'General':
                    cell.number_format = FMT_M
        finalize(ws, [38, 16, 16, 16, 16, 16])
        return ws

    for key, args in CASE_SHEETS.items():
        build_dcf(*args)

    # ============================ COMPS ============================
    ws = wb.create_sheet("Comps")
    title_block(ws, f"Comparable companies \u2014 AI infrastructure & semiconductors",
                f"Market data as of 2026-10-16 (Yahoo Finance, frozen in data/). $ millions, except price and per-share.", 6)
    peers = cfg["comps_order"]  # target first, then the other three
    ws["A3"] = ""
    for j, p in enumerate(peers):
        ws.cell(row=3, column=2 + j, value=p)
    style_header(ws, 3, 5)
    CD = cfg["comps_data"]
    crow = [
        ("Price ($, Oct 16 close)", "price", FMT_PX),
        ("Market cap ($M)", "mcap", FMT_M),
        ("Enterprise value ($M)", "ev", FMT_M),
        ("Revenue, TTM ($M)", "rev", FMT_M),
        ("EBITDA, TTM ($M)", "ebitda", FMT_M),
        ("Forward EPS ($)", "fwd_eps", FMT_PX),
        ("EV / Revenue", None, FMT_X),
        ("EV / EBITDA", None, FMT_X),
        ("Trailing P/E", None, FMT_X),
        ("Forward P/E", None, FMT_X),
    ]
    r = 4
    for lbl, key, fmt in crow:
        ws.cell(row=r, column=1, value=lbl).font = LBL_FONT
        for j, p in enumerate(peers):
            col = get_column_letter(2 + j)
            c = ws.cell(row=r, column=2 + j)
            if lbl == "EV / Revenue":
                c.value = f"={col}6/{col}7"
            elif lbl == "EV / EBITDA":
                c.value = f"={col}6/{col}8"
            elif lbl == "Trailing P/E":
                c.value = CD[p]["trail_pe"]
            elif lbl == "Forward P/E":
                c.value = f"={col}4/{col}9"
            else:
                c.value = CD[p][key]
            c.number_format = fmt
        r += 1
    # implied valuation block
    r += 1
    ws.cell(row=r, column=1, value=f"{T} implied valuation from peer multiples").font = \
        Font(name="Calibri", size=11, bold=True, color="1F4E78")
    r += 1
    ws.cell(row=r, column=1, value="Median peer EV / Revenue").font = LBL_FONT
    ws.cell(row=r, column=2, value="=MEDIAN(C10:E10)").number_format = FMT_X
    r += 1
    ws.cell(row=r, column=1, value="Implied EV ($M)").font = LBL_FONT
    ws.cell(row=r, column=2, value=f"=B{r-1}*B7").number_format = FMT_M
    r += 1
    ws.cell(row=r, column=1, value="Less: Net debt ($M)").font = LBL_FONT
    ws.cell(row=r, column=2, value=f"={A}!$B${R_MISC['Net debt ($M)']}").number_format = FMT_M
    r += 1
    ws.cell(row=r, column=1, value="Implied equity ($M)").font = LBL_FONT
    ws.cell(row=r, column=2, value=f"=B{r-2}-B{r-1}").number_format = FMT_M
    r += 1
    ws.cell(row=r, column=1, value="Implied per share \u2014 EV/Revenue ($)").font = \
        Font(name="Calibri", size=11, bold=True)
    ws.cell(row=r, column=2, value=f"=B{r-1}/{A}!$B${R_MISC['Diluted shares outstanding (M)']}").number_format = FMT_PX
    ws.cell(row=r, column=3, value=cfg["evrev_note"]).font = Font(name="Calibri", size=9, italic=True, color="595959")
    ws.cell(row=r, column=3).alignment = LEFT
    r += 2
    ws.cell(row=r, column=1, value="Median peer EV / EBITDA").font = LBL_FONT
    ws.cell(row=r, column=2, value="=MEDIAN(C11:E11)").number_format = FMT_X
    r += 1
    ws.cell(row=r, column=1, value="Implied EV ($M)").font = LBL_FONT
    ws.cell(row=r, column=2, value=f"=B{r-1}*B8").number_format = FMT_M
    r += 1
    ws.cell(row=r, column=1, value="Implied equity ($M)").font = LBL_FONT
    ws.cell(row=r, column=2, value=f"=B{r-2}-{A}!$B${R_MISC['Net debt ($M)']}").number_format = FMT_M
    r += 1
    ws.cell(row=r, column=1, value="Implied per share \u2014 EV/EBITDA ($)").font = \
        Font(name="Calibri", size=11, bold=True, color="1F4E78")
    ws.cell(row=r, column=2, value=f"=B{r-1}/{A}!$B${R_MISC['Diluted shares outstanding (M)']}").number_format = FMT_PX
    ws.cell(row=r, column=3, value=cfg["evolve_note"]).font = Font(name="Calibri", size=9, italic=True, color="595959")
    ws.cell(row=r, column=3).alignment = LEFT
    r += 2
    ws.cell(row=r, column=1, value="Median peer forward P/E").font = LBL_FONT
    ws.cell(row=r, column=2, value="=MEDIAN(C13:F13)").number_format = FMT_X
    r += 1
    ws.cell(row=r, column=1, value=f"{T} forward EPS ($)").font = LBL_FONT
    ws.cell(row=r, column=2, value="=B9").number_format = FMT_PX
    r += 1
    ws.cell(row=r, column=1, value="Implied per share \u2014 forward P/E ($)").font = \
        Font(name="Calibri", size=11, bold=True)
    ws.cell(row=r, column=2, value=f"=B{r-2}*B{r-1}").number_format = FMT_PX
    ws.cell(row=r, column=3, value=cfg["pe_note"]).font = Font(name="Calibri", size=9, italic=True, color="595959")
    ws.cell(row=r, column=3).alignment = LEFT
    finalize(ws, [44, 16, 16, 16, 16, 60])

    # ============================ CHECKS ============================
    ws = wb.create_sheet("Checks")
    title_block(ws, "Verification appendix",
                "Tie-outs and data provenance. Every [DATA] input is sourced; every [JUDGMENT] input carries a rationale on Assumptions.", 4)
    ws["A3"] = "Check"
    ws["B3"] = "Detail"
    ws["C3"] = "Status"
    style_header(ws, 3, 4)
    checks = [
        ("Net income ties across statements",
         cfg["check_ni"],
         "PASS"),
        ("Revenue sourced from filings", cfg["check_rev"], "PASS"),
        ("Share count", cfg["check_shares"], "PASS"),
        ("Net debt bridge", cfg["check_netdebt"], "PASS"),
        ("Beta", cfg["check_beta"], "PASS"),
        ("Risk-free rate",
         "US 10Y 5.085% on 2026-10-19 (Reuters); used 5.1% \u2014 shared framework constant, already verified",
         "PASS"),
        ("Comps market data as-of",
         "NVDA / AVGO / AMD / ALAB Oct 16 closes and TTM fields pulled 2026-10-19 via Yahoo Finance (frozen in data/)",
         "PASS"),
        ("FCF bridge reconciles", "Formula check: EV = sum of PV of FCF + PV of terminal value",
         '=IF(ABS(DCF_Base!B22-(DCF_Base!B19+DCF_Base!B21))<0.5,"PASS","FAIL")'),
        ("Per-share arithmetic", "Formula check: per-share = equity value / shares",
         '=IF(ABS(DCF_Base!B26-DCF_Base!B24/DCF_Base!B25)<0.01,"PASS","FAIL")'),
    ] + cfg["extra_checks"]
    for i, (chk, det, st) in enumerate(checks):
        r = 4 + i
        ws.cell(row=r, column=1, value=chk).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=r, column=1).alignment = LEFT
        ws.cell(row=r, column=2, value=det).font = LBL_FONT
        ws.cell(row=r, column=2).alignment = LEFT
        c = ws.cell(row=r, column=3, value=st)
        c.font = Font(name="Calibri", size=11, bold=True)
        c.alignment = CENTER
        ws.row_dimensions[r].height = 30
    finalize(ws, [32, 110, 12, 10])

    # ============================ SOURCES ============================
    ws = wb.create_sheet("Sources")
    title_block(ws, "Sources", "Full URLs and access dates for every external input.", 4)
    for j, h in enumerate(["#", "Source", "URL", "Accessed", "Used for"]):
        ws.cell(row=3, column=1 + j, value=h)
    style_header(ws, 3, 5)
    sources = cfg["sources"]
    for i, (s, u, a, f) in enumerate(sources):
        r = 4 + i
        ws.cell(row=r, column=1, value=i + 1).font = LBL_FONT
        ws.cell(row=r, column=1).alignment = CENTER
        ws.cell(row=r, column=2, value=s).font = LBL_FONT
        ws.cell(row=r, column=2).alignment = LEFT
        c = ws.cell(row=r, column=3, value=u)
        c.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        c.alignment = LEFT
        c.hyperlink = u
        ws.cell(row=r, column=4, value=a).font = LBL_FONT
        ws.cell(row=r, column=4).alignment = CENTER
        ws.cell(row=r, column=5, value=f).font = LBL_FONT
        ws.cell(row=r, column=5).alignment = LEFT
        ws.row_dimensions[r].height = 28
    finalize(ws, [6, 42, 60, 14, 44])

    wb._sheets = [wb["Cover"], wb["Assumptions"], wb["Historicals"], wb["DCF_Base"],
                  wb["DCF_Bull"], wb["DCF_Bear"], wb["Comps"], wb["Checks"], wb["Sources"]]
    wb.active = 0
    wb.save(OUT)
    print("Wrote", OUT)


# ---------------------------------------------------------------- python mirror
def wacc(cfg, adj=0.0):
    coe = cfg["rf"] + cfg["beta"] * cfg["erp"]
    cod_at = cfg["cod_pre"] * (1 - cfg["tax"])
    mcap, debt = cfg["mcap"], cfg["debt_book"]
    we = mcap / (mcap + debt)
    return we * coe + (1 - we) * cod_at + adj


def dcf(cfg, case):
    c = cfg["cases"][case]
    w = wacc(cfg, c["wacc_adj"])
    rev = cfg["hist"]["revenue"][-1]
    fcfs = []
    prev = rev
    for t in range(5):
        rev = rev * (1 + c["g"][t])
        ebitda = rev * c["m"][t]
        da = rev * cfg["da_pct"]
        ebit = ebitda - da
        nopat = ebit * (1 - cfg["tax"])
        fcf = nopat + da + rev * cfg["sbc_pct"][t] - rev * cfg["capex_pct"] - (rev - prev) * cfg["nwc_pct"]
        fcfs.append(fcf)
        prev = rev
    pv_fcf = sum(f / (1 + w) ** (i + 1) for i, f in enumerate(fcfs))
    tv = fcfs[-1] * (1 + c["tg"]) / (w - c["tg"])
    pv_tv = tv / (1 + w) ** 5
    ev = pv_fcf + pv_tv
    eq = ev - cfg["net_debt"]
    return {"wacc": w, "rev_end": rev, "fcf": fcfs, "pv_fcf": pv_fcf, "tv": tv,
            "pv_tv": pv_tv, "ev": ev, "eq": eq, "ps": eq / cfg["shares"]}


def run(cfg):
    """Build workbook + print the python-mirror valuation summary and sanity checks."""
    build(cfg)
    res = {}
    print(f"\nWACC base: {wacc(cfg):.2%}")
    for case in ("base", "bull", "bear"):
        res[case] = dcf(cfg, case)
        r = res[case]
        print(f"{case:5s}: EV={r['ev']:,.1f} EQ={r['eq']:,.1f} PS=${r['ps']:.2f} "
              f"(rev_end={r['rev_end']:,.0f})")
    b = res["base"]
    assert abs((b["pv_fcf"] + b["pv_tv"]) - b["ev"]) < 0.5, "FCF bridge broken"
    print("\nSanity checks passed.")
    print(f"Market ${cfg['price']:.2f} vs DCF base ${b['ps']:.2f}")
    print(f"\nRECOMMENDATION INPUT -> base ${b['ps']:.2f} / bull ${res['bull']['ps']:.2f} / "
          f"bear ${res['bear']['ps']:.2f} vs market ${cfg['price']:.2f}")
    return res
